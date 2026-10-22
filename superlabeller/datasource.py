"""Loading records for printing from CSV, Excel, SQLite and ODBC databases.

A data source is a dict saved with the design, e.g.
  {"kind": "csv",    "path": "items.csv", "delimiter": "auto", "encoding": "auto", "header": True}
  {"kind": "xlsx",   "path": "items.xlsx", "sheet": "Sheet1", "header": True}
  {"kind": "sqlite", "path": "shop.db", "query": "SELECT * FROM products"}
  {"kind": "odbc",   "connection": "Driver={SQL Server};Server=...;", "query": "SELECT ..."}

load() returns (column names, list of row dicts). Values keep their type
(numbers, dates) so placeholders like {Price:.2f} can format them.
"""
import csv
import io
import os
import sqlite3

KINDS = [
    ("csv", "CSV / text file"),
    ("xlsx", "Excel workbook (.xlsx)"),
    ("sqlite", "SQLite database"),
    ("odbc", "ODBC database (SQL Server, Access, MySQL, PostgreSQL...)"),
]
MAX_ROWS = 200000


class DataError(Exception):
    pass


def describe(src):
    if not src:
        return "No data source"
    kind = src.get("kind")
    if kind in ("csv", "xlsx", "sqlite"):
        name = os.path.basename(src.get("path", ""))
        if kind == "xlsx" and src.get("sheet"):
            name += f" [{src['sheet']}]"
        return name
    if kind == "odbc":
        if src.get("connection"):
            where = "custom connection"
        else:
            where = src.get("dsn") or "/".join(p for p in (src.get("server"), src.get("database")) if p)
        table = f" - {src['table']}" if src.get("table") else ""
        return f"{where or 'database'}{table}"
    return "?"


def sqlite_columns(path, table):
    con = sqlite3.connect(path)
    try:
        return [r[1] for r in con.execute(f'PRAGMA table_info("{table}")')]
    finally:
        con.close()


def _unique_columns(names):
    out, seen = [], set()
    for i, n in enumerate(names):
        n = str(n).strip() if n not in (None, "") else f"Column{i + 1}"
        base, k = n, 2
        while n in seen:
            n = f"{base}_{k}"
            k += 1
        seen.add(n)
        out.append(n)
    return out


def _table(rows, header):
    rows = [r for r in rows if any(v not in (None, "") and str(v).strip() for v in r)]
    if not rows:
        return [], []
    width = max(len(r) for r in rows)
    head = list(rows[0]) + [None] * (width - len(rows[0])) if header else [None] * width
    cols = _unique_columns(head)
    body = rows[1:] if header else rows
    return cols, [{c: (r[i] if i < len(r) and r[i] is not None else "") for i, c in enumerate(cols)}
                  for r in body[:MAX_ROWS]]


def _need_file(path):
    if not path or not os.path.isfile(path):
        raise DataError(f"File not found: {path or '(none)'}")


# ------------------------------------------------------------------ CSV
def _load_csv(src):
    path = src.get("path", "")
    _need_file(path)
    with open(path, "rb") as f:
        raw = f.read()
    enc = src.get("encoding", "auto")
    if enc == "auto":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252", "replace")
    else:
        text = raw.decode(enc, "replace")
    delim = src.get("delimiter", "auto")
    if delim == "auto":
        try:
            delim = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    return _table(list(csv.reader(io.StringIO(text), delimiter=delim)), src.get("header", True))


# ------------------------------------------------------------------ Excel
def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise DataError("Reading Excel files needs the 'openpyxl' package (pip install openpyxl).")


def sheet_names(path):
    wb = _openpyxl().load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _load_xlsx(src):
    path = src.get("path", "")
    _need_file(path)
    if path.lower().endswith(".xls"):
        raise DataError("Old .xls files are not supported. Save the file as .xlsx (or CSV) in Excel.")
    try:
        wb = _openpyxl().load_workbook(path, read_only=True, data_only=True)
    except DataError:
        raise
    except Exception as e:
        raise DataError(f"Could not open workbook: {e}")
    try:
        sheet = src.get("sheet")
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return _table(rows, src.get("header", True))


# ------------------------------------------------------------------ SQL
def _run_query(con, query):
    if not query.strip():
        raise DataError("Enter a query, e.g. SELECT * FROM products")
    try:
        cur = con.cursor()
        cur.execute(query)
        if not cur.description:
            raise DataError("The query did not return any columns.")
        cols = _unique_columns([d[0] for d in cur.description])
        rows = cur.fetchmany(MAX_ROWS)
    except DataError:
        raise
    except Exception as e:
        raise DataError(f"Query failed: {e}")
    return cols, [dict(zip(cols, r)) for r in rows]


def sqlite_tables(path):
    _need_file(path)
    con = sqlite3.connect(path)
    try:
        return [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table','view') "
            "AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    finally:
        con.close()


def _load_sqlite(src):
    path = src.get("path", "")
    _need_file(path)
    con = sqlite3.connect(path)
    try:
        return _run_query(con, src.get("query", ""))
    finally:
        con.close()


def _pyodbc():
    try:
        import pyodbc
        return pyodbc
    except ImportError:
        raise DataError("Database access needs the 'pyodbc' package (pip install pyodbc).")


def odbc_drivers():
    try:
        return list(_pyodbc().drivers())
    except DataError:
        return []


def driver_family(driver):
    d = (driver or "").lower()
    for key, fam in (("sql server", "mssql"), ("mysql", "mysql"), ("maria", "mysql"), ("postgre", "postgres"),
                     ("access", "access"), ("oracle", "oracle"), ("sqlite", "sqlite")):
        if key in d:
            return fam
    return "generic"


def _q(v):
    """Quote an ODBC connection-string value if needed."""
    v = str(v)
    return "{" + v.replace("}", "}}") + "}" if any(c in v for c in ";{}=") else v


def needs_password(src):
    return (src.get("kind") == "odbc" and not src.get("connection") and not src.get("trusted")
            and bool(src.get("user")) and not src.get("password"))


def odbc_connection_string(src):
    """Build an ODBC connection string from the credential fields (or use the custom one)."""
    if src.get("connection"):
        return src["connection"]
    d = src.get("driver", "")
    fam = driver_family(d)
    server, port, db = src.get("server", ""), src.get("port", ""), src.get("database", "")
    user, pw = src.get("user", ""), src.get("password", "")
    parts = [f"DSN={_q(src['dsn'])}"] if src.get("dsn") else [f"Driver={{{d}}}"]
    if fam == "access":
        parts.append(f"DBQ={_q(db)}")
    elif fam == "mssql":
        parts.append(f"Server={_q(server + (',' + str(port) if port else ''))}")
        if db:
            parts.append(f"Database={_q(db)}")
        if "18" in d or "17" in d:
            parts.append("TrustServerCertificate=yes")
    elif server or db:
        if server:
            parts.append(f"Server={_q(server)}")
        if port:
            parts.append(f"Port={port}")
        if db:
            parts.append(f"Database={_q(db)}")
    if src.get("trusted") and fam == "mssql":
        parts.append("Trusted_Connection=yes")
    else:
        if user:
            parts.append(f"UID={_q(user)}")
        if pw:
            parts.append(f"PWD={_q(pw)}")
    return ";".join(parts) + ";"


def _odbc_connect(src):
    if not (src.get("connection") or src.get("driver") or src.get("dsn")):
        raise DataError("Choose a database driver.")
    try:
        return _pyodbc().connect(odbc_connection_string(src), timeout=10, autocommit=True)
    except DataError:
        raise
    except Exception as e:
        raise DataError(f"Could not connect: {e}")


def odbc_tables(src):
    con = _odbc_connect(src)
    try:
        out = []
        for r in con.cursor().tables():
            if r.table_type not in ("TABLE", "VIEW"):
                continue
            if (r.table_schem or "").lower() in ("sys", "information_schema", "pg_catalog"):
                continue
            out.append(f"{r.table_schem}.{r.table_name}" if r.table_schem else r.table_name)
        return out
    except Exception as e:
        raise DataError(f"Could not list tables: {e}")
    finally:
        con.close()


def odbc_columns(src, table):
    con = _odbc_connect(src)
    try:
        schema, _, name = table.rpartition(".")
        cur = con.cursor()
        rows = cur.columns(table=name, schema=schema or None).fetchall()
        return [r.column_name for r in rows]
    except Exception as e:
        raise DataError(f"Could not read columns: {e}")
    finally:
        con.close()


def quote_table(src, table):
    """SELECT-able table reference with the right identifier quoting."""
    kind = src.get("kind")
    fam = driver_family(src.get("driver", "")) if kind == "odbc" else "sqlite"
    if fam in ("mssql", "access"):
        l, r = "[", "]"
    elif fam == "mysql":
        l, r = "`", "`"
    else:
        l, r = '"', '"'
    return ".".join(f"{l}{p}{r}" for p in table.split(".")) if fam != "access" else f"[{table}]"


def _load_odbc(src):
    con = _odbc_connect(src)
    try:
        return _run_query(con, src.get("query", ""))
    finally:
        con.close()


def load(src):
    loaders = {"csv": _load_csv, "xlsx": _load_xlsx, "sqlite": _load_sqlite, "odbc": _load_odbc}
    if not src or src.get("kind") not in loaders:
        raise DataError("No data source selected.")
    return loaders[src["kind"]](src)
